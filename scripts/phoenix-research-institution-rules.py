#!/usr/bin/env python3
"""Extract cited *candidate* academic layout rules from a user-authorized local file.

Text interpretation is conservative: nothing here certifies institutional compliance.
Reads PDF/DOCX/TXT/MD only, never executes macros, external links, or embedded objects.
"""
import json
import re
import sys
from pathlib import Path

MAX_BYTES = 15 * 1024 * 1024
MAX_PAGES = 200
RULES = [
    ("margins", re.compile(r"\b(?:m[aá]rgen(?:es)?|margins?)\b", re.I)),
    ("line_spacing", re.compile(r"\b(?:interlineado|espaciado entre l[ií]neas|double.spacing|line.spacing)\b", re.I)),
    ("font", re.compile(r"\b(?:tipograf[ií]a|fuente(?:\s+(?:de\s+)?letra)?|tipo\s+de\s+letra|Arial|Calibri|Times New Roman|Georgia)\b", re.I)),
    ("page_count", re.compile(r"\b(?:p[aá]ginas?|extensi[oó]n|l[ií]mite de palabras|word.limit)\b", re.I)),
    ("abstract", re.compile(r"\b(?:abstract|resumen|palabras clave|keywords)\b", re.I)),
    ("sections", re.compile(r"\b(?:cap[ií]tulos?|marco te[oó]rico|resultados|metodolog[ií]a|conclusiones|secciones)\b", re.I)),
    ("citations", re.compile(r"\b(?:APA|Vancouver|AMA|citas?|referencias?|bibliograf[ií]a)\b", re.I)),
    ("cover", re.compile(r"\b(?:portada|logotipo|escudo|matr[ií]cula|participantes?|asesor(?:a)?|tutor(?:a)?)\b", re.I)),
]
def get_pages(path):
    extension = path.suffix.lower()
    if extension == ".pdf":
        import fitz
        document = fitz.open(str(path))
        if document.page_count > MAX_PAGES:
            raise ValueError("institutional PDF exceeds 200 pages")
        return [(f"página {number+1}", page.get_text("text")) for number, page in enumerate(document)]
    if extension == ".docx":
        from docx import Document
        word = Document(path)
        units = []
        for index, paragraph in enumerate(word.paragraphs):
            if paragraph.text.strip():
                units.append((f"párrafo {index+1}", paragraph.text))
        for index, table in enumerate(word.tables):
            for row in table.rows:
                units.append((f"tabla {index+1}", " | ".join(cell.text for cell in row.cells)))
        return units[:5000]
    if extension in (".txt", ".md"):
        return [(f"línea {index+1}", line) for index, line in
                enumerate(path.read_text(encoding="utf-8-sig").splitlines()[:8000])]
    raise ValueError("only .pdf, .docx, .txt and .md institutional guides are supported")

def main():
    request = json.loads(sys.stdin.buffer.read(8192).decode("utf-8"))
    supplied = request.get("path")
    if not isinstance(supplied, str) or len(supplied) > 2000:
        raise ValueError("path must be user supplied")
    path = Path(supplied).expanduser().resolve(strict=True)
    if not path.is_file() or path.stat().st_size > MAX_BYTES:
        raise ValueError("guide missing, not a file, or exceeds 15 MB")
    pages = get_pages(path)
    findings = []
    text_pages = 0
    for locator, text in pages:
        if text.strip(): text_pages += 1
        for line in text.splitlines():
            snippet = " ".join(line.split())
            if len(snippet) < 10: continue
            for kind, regex in RULES:
                if regex.search(snippet):
                    findings.append({
                        "kind": kind, "source_locator": locator, "source_excerpt": snippet[:350],
                        "review_required": True,
                    })
                    if len(findings) >= 150: break
            if len(findings) >= 150: break
        if len(findings) >= 150: break
    warnings = []
    if not text_pages: warnings.append("No se extrajo texto: quizá es un PDF escaneado. Se requiere lectura visual/OCR aprobada.")
    if not findings: warnings.append("No se detectaron reglas claras mediante la búsqueda textual; revisar el documento.")
    result = {
        "state": "RULES_FOR_REVIEW", "file_name": path.name,
        "inspected_units": len(pages), "text_units": text_pages,
        "candidates": findings, "warnings": warnings,
        "disclaimer": "Extracción textual heurística con localizadores; no se afirma conformidad institucional hasta revisar y aplicar las reglas.",
    }
    print(json.dumps(result, ensure_ascii=True), flush=True)

if __name__ == "__main__":
    try: main()
    except Exception as error:
        sys.stderr.write("RESEARCH_RULES_FAILED: " + str(error) + "\n")
        sys.exit(2)
