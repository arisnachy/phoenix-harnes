#!/usr/bin/env python3
"""Exercise institutional-guide extraction without real user documents."""
import json
import subprocess
import sys
import tempfile
from pathlib import Path
from docx import Document
import pymupdf as fitz

script = Path(__file__).with_name("phoenix-research-institution-rules.py")
with tempfile.TemporaryDirectory(prefix="research-rules-ci-") as root:
    folder = Path(root)
    word = Document()
    word.add_paragraph("Márgenes: 2.54 cm en cada lado.")
    word.add_paragraph("Fuente Arial 11 puntos; interlineado doble.")
    word.add_paragraph("Portada institucional y referencias bibliográficas APA 7.")
    docx_path = folder/"manual.docx"
    word.save(docx_path)
    pdf_path = folder/"manual.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((55, 60), "Margins: 1 inch, double-spacing. APA reference format.")
    pdf.save(pdf_path)
    for path in (docx_path, pdf_path):
        process = subprocess.run([sys.executable, str(script)],
                                 input=json.dumps({"path": str(path)}).encode("utf-8"),
                                 capture_output=True, check=True, timeout=40)
        assert process.stdout.strip(), f"Extractor produced empty output: {path} / stderr={process.stderr.decode('utf-8', errors='replace')!r}"
        result = json.loads(process.stdout.decode("utf-8"))
        assert result["state"] == "RULES_FOR_REVIEW"
        assert result["candidates"] and all(row["source_locator"] and row["source_excerpt"]
                                            and row["review_required"] for row in result["candidates"])
        assert result["text_units"] >= 1
    refused = subprocess.run([sys.executable, str(script)],
                             input=json.dumps({"path": str(folder/"missing.docx")}).encode("utf-8"),
                             capture_output=True, timeout=25)
    assert refused.returncode != 0
print("RESEARCH_INSTITUTION_RULES_SMOKE_PASS")
