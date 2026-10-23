#!/usr/bin/env python3
"""Cross-platform real DOCX, PDF and HTML smoke test for Research Studio."""
import json
import subprocess
import sys
from pathlib import Path
from zipfile import ZipFile
from PIL import Image
from docx import Document
import fitz

script = Path(__file__).with_name("phoenix-research-publisher.py")
logo = Path(__file__).with_name(".phoenix-research-test-logo.png")
Image.new("RGB", (160, 100), "#c67a3c").save(logo)
try:
    report = {
        "title": "Informe de prueba de Phoenix Research Studio",
        "style": "apa",
        "presentation": "academic",
        "document_type": "brief", "quality_profile": "educational_brief", "user_requested_brief": True,
        "keywords": ["metodología", "fuentes", "validación"],
        "paper": "A4",
        "cover": {"institution": "Institución de prueba", "faculty": "Facultad de Ciencias",
                  "author": "Investigador de prueba", "date": "2026-10-17",
                  "logo_path": str(logo)},
        "sections": [
            {"heading": "Resumen", "text": "Dos fuentes de prueba [@source1; @source2] sin marcadores literales."},
            {"heading": "Metodología", "text": "**Implicaciones para investigación.** Hay *cursivas* y `código` sin etiquetas crudas, con citas [@source1]."},
            {"heading": "Limitaciones", "text": "Esta prueba no valida la evidencia externa [@source1]."},
        ],
        "references": [
            {"key": "source1", "authors": "Apellido, A.", "year": "2026",
             "title": "Referencia sintética de prueba", "journal": "Revista de prueba",
             "url": "https://example.org/phoenix-research-ci"},
            {"key": "source2", "authors": "Otro, B.", "year": "2024",
             "title": "Otra referencia sintética de prueba", "journal": "Revista de prueba",
             "url": "https://example.org/phoenix-research-ci-2"},
        ],
        "tables": [
            {"title": "Comparación de prueba", "columns": ["Grupo", "Tamaño", "Resultado"],
             "rows": [["A", 12, 4], ["B", 18, 9]], "source_key": "source1",
             "note": "Valores sintéticos únicamente para pruebas de formato"}
        ],
        "charts": [
            {"title": "Gráfico sintético", "labels": ["A", "B", "C"],
             "values": [4, 9, 3], "source_key": "source1"},
            {"title": "Forest plot sintético", "type": "forest", "labels": ["Estudio A", "Estudio B"],
             "values": [1.2, 0.8], "ci_low": [0.9, 0.6], "ci_high": [1.5, 1.1],
             "null_value": 1, "method": "Efectos de prueba con IC sintéticos",
             "source_key": "source1"},
        ],
    }
    for style, presentation in (("apa", "academic"), ("vancouver", "hybrid"), ("apa7", "editorial")):
        report["style"] = style
        report["presentation"] = presentation
        process = subprocess.run([sys.executable, str(script)], input=json.dumps(report).encode("utf-8"),
                                 capture_output=True, timeout=90)
        assert process.returncode == 0, process.stderr.decode("utf-8", errors="replace")
        receipt = json.loads(process.stdout)
        assert receipt["state"] == "REPORT_READY", receipt
        assert Path(receipt["project_path"]).is_file()
        assert json.loads(Path(receipt["project_path"]).read_text(encoding="utf-8"))["style"] == style
        assert receipt["source_verification"] == "metadata_user_supplied_not_network_verified"
        paths = {k: Path(v) for k, v in receipt["files"].items()}
        assert all(p.is_file() and p.stat().st_size > 500 for p in paths.values())
        assert paths["pdf"].read_bytes().startswith(b"%PDF-")
        with fitz.open(paths["pdf"]) as pdf:
            assert pdf.page_count >= 2
            report_text = "\n".join(page.get_text() for page in pdf)
            assert "Referencias" in report_text
            assert "[@" not in report_text
            assert "[1,2]" in report_text if style == "vancouver" else "[@" not in report_text
            assert "Tabla 1" in report_text and "Figura 2" in report_text
            assert "Implicaciones para investigación." in report_text
            assert "**" not in report_text and "*cursivas*" not in report_text
            for page in pdf:
                assert page.rect.width >= 500 and page.rect.height >= 700
                for block in page.get_text("blocks"):
                    if len(block) > 6 and block[6] != 0:
                        continue
                    x0, y0, x1, y1 = block[:4]
                    assert -1 <= x0 < x1 <= page.rect.width + 1
                    assert -1 <= y0 < y1 <= page.rect.height + 1
        html_contents = paths["html"].read_text(encoding="utf-8")
        assert "Referencias" in html_contents
        assert "[@" not in html_contents
        assert "<table>" in html_contents and "Comparación de prueba" in html_contents
        assert "<strong>Implicaciones para investigación.</strong>" in html_contents
        assert "<em>cursivas</em>" in html_contents and "<code>código</code>" in html_contents
        assert "**" not in html_contents
        assert "Figura 2" in html_contents and "Forest plot sintético" in html_contents and "alt=" in html_contents
        assert receipt["presentation"] == presentation
        assert receipt["quality"]["submission_certified"] is False
        with ZipFile(paths["docx"]) as archive:
            assert archive.testzip() is None
            assert "word/document.xml" in archive.namelist()
            assert any(name.startswith("word/media/") for name in archive.namelist())
            xml = archive.read("word/document.xml").decode("utf-8")
            assert "Referencias" in xml and "Metodología" in xml
            assert "[@" not in xml
            assert "Comparación de prueba" in xml
            assert "**" not in xml
        assert any(run.bold and "Implicaciones para investigación." in run.text
                   for paragraph in Document(paths["docx"]).paragraphs for run in paragraph.runs)
        assert any(run.italic and "cursivas" in run.text
                   for paragraph in Document(paths["docx"]).paragraphs for run in paragraph.runs)
            assert ("(Apellido, 2026)" in xml) if style in ("apa", "apa7") else ("[1]" in xml)
        print("OK", style, paths["pdf"], paths["docx"], paths["html"])
    report["document_type"] = "thesis"
    incomplete = subprocess.run([sys.executable, str(script)], input=json.dumps(report).encode("utf-8"),
                                capture_output=True, timeout=30)
    assert incomplete.returncode != 0
    assert b"requires missing sections" in incomplete.stderr
    report["document_type"] = "systematic_review"
    not_systematic = subprocess.run([sys.executable, str(script)], input=json.dumps(report).encode("utf-8"),
                                     capture_output=True, timeout=30)
    assert not_systematic.returncode != 0
    thesis = dict(report)
    thesis["style"] = "apa7"
    thesis["presentation"] = "hybrid"
    thesis["document_type"] = "brief"
    thesis["table_of_contents"] = True
    thesis["cover"] = {
        "institution": "Universidad de prueba", "faculty": "Ciencias de la Salud",
        "department": "Investigación", "degree": "Maestría", "advisor": "Tutora de prueba",
        "course": "Seminario", "course_code": "INV-100", "city": "Santiago",
        "country": "República Dominicana", "date": "2026-10-18",
        "authors": [{"name": "Autora Uno", "student_id": "MA-001"},
                    {"name": "Autor Dos", "student_id": "MA-002"}],
    }
    thesis["front_matter"] = {
        "originality": "Los datos de este ejemplo son sintéticos.",
        "dedication": "Dedicatoria de prueba.",
        "acknowledgements": "Agradecimientos de prueba.",
        "abstract_en": "English abstract for a synthetic publishing test.",
    }
    thesis["page_plan"] = {"target_pages": 80, "scope": "all"}
    thesis["declarations"] = {
        "ethics": "No hubo participantes humanos en este ejemplo sintético.",
        "funding": "Sin financiación; ejemplo de prueba.",
        "conflicts": "Sin conflictos declarados en esta prueba.",
    }
    thesis["appendices"] = [
        {"title": "Instrumento sintético", "text": "Apéndice con fuente de prueba [@source1]."}
    ]
    thesis["sections"] = [
        {"kind": kind, "heading": heading, "text": "Ejemplo documentado solo para probar maquetación [@source1]."}
        for kind, heading in (
            ("abstract", "Resumen"), ("introduction", "Introducción"),
            ("problem", "Planteamiento del problema"), ("objectives", "Objetivos"),
            ("framework", "Marco teórico"), ("methods", "Metodología"),
            ("results", "Resultados"), ("discussion", "Discusión"),
            ("conclusions", "Conclusiones"), ("recommendations", "Recomendaciones"),
        )
    ]
    process = subprocess.run([sys.executable, str(script)], input=json.dumps(thesis).encode("utf-8"),
                             capture_output=True, check=True, timeout=90)
    receipt = json.loads(process.stdout)
    with fitz.open(receipt["files"]["pdf"]) as actual:
        assert receipt["pagination"]["pdf_pages"] == actual.page_count
        assert actual.page_count < 80
        pages_text = "\\n".join(page.get_text() for page in actual)
        assert "Índice general" in pages_text and "Autora Uno" in pages_text
        assert "Anexos y apéndices" in pages_text and "Declaraciones de investigación" in pages_text
    assert receipt["state"] == "REPORT_NEEDS_REVISION"
    assert receipt["pagination"]["target_met"] is False
    assert receipt["quality"]["pagination_review_required"] is True
    assert receipt["warnings"] and "80" in receipt["warnings"][0]
    with ZipFile(receipt["files"]["docx"]) as archive:
        word = archive.read("word/document.xml").decode("utf-8")
        assert "MA-001" in word and "MA-002" in word
        assert "Dedicatoria" in word and "Índice de tablas" in word
        assert "Instrumento sintético" in word and "Conflictos de interés" in word
        assert 'TOC ' in word and 'w:fldSimple' in word
    journal = dict(report)
    journal["style"] = "ama"
    journal["document_type"] = "brief"
    journal["institutional_rules"] = {
        "reviewed_by_user": True, "margin_cm": 2.54,
        "font_family": "Arial", "line_spacing": 2,
    }
    journal["front_matter"] = {"key_points": [
        "Pregunta: publicación con evidencia sintética.",
        "Hallazgos: el diseño genera archivos comprobables.",
        "Significado: el autor debe verificar resultados verdaderos.",
    ]}
    journal["journal_requirements"] = {
        "journal": "jama", "max_words": 1000, "max_abstract_words": 350,
        "max_references": 5, "max_tables_figures": 5, "require_key_points": True,
    }
    published = subprocess.run([sys.executable, str(script)],
                               input=json.dumps(journal, ensure_ascii=False).encode("utf-8"),
                               capture_output=True, check=True, timeout=90)
    result = json.loads(published.stdout.decode("utf-8"))
    assert result["state"] == "REPORT_READY"
    with ZipFile(result["files"]["docx"]) as document:
        xml = document.read("word/document.xml").decode("utf-8")
        assert "Key Points" in xml
    with fitz.open(result["files"]["pdf"]) as document:
        text = "\n".join(page.get_text() for page in document)
        assert "Key Points" in text
    oversized = dict(journal)
    oversized["journal_requirements"] = {**journal["journal_requirements"], "max_words": 1}
    rejected = subprocess.run([sys.executable, str(script)],
                              input=json.dumps(oversized).encode("utf-8"),
                              capture_output=True, timeout=30)
    assert rejected.returncode != 0
    unapproved = dict(journal)
    unapproved["institutional_rules"] = {"margin_cm": 3}
    rejected = subprocess.run([sys.executable, str(script)],
                              input=json.dumps(unapproved).encode("utf-8"),
                              capture_output=True, timeout=30)
    assert rejected.returncode != 0
    shallow = dict(report)
    shallow["quality_profile"] = "professional"
    shallow["document_type"] = "narrative_review"
    shallow["sections"] = [
        {"kind": kind, "heading": label,
         "text": "Este párrafo no contiene suficiente investigación [@source1; @source2]."}
        for kind, label in (("abstract", "Resumen"), ("introduction", "Introducción"),
            ("methods", "Método"), ("results", "Resultados"),
            ("discussion", "Discusión"), ("limitations", "Limitaciones"),
            ("conclusions", "Conclusiones"))
    ]
    refused = subprocess.run([sys.executable, str(script)],
                             input=json.dumps(shallow).encode("utf-8"),
                             capture_output=True, timeout=30)
    assert refused.returncode != 0, "A three-page shallow review must not publish."
    assert b"RESEARCH_QUALITY_GATE" in refused.stderr, refused.stderr
    # Positive integration: a substantial hybrid research report must really render,
    # not merely satisfy a validator with fake filenames. All data here are synthetic.
    developed = dict(report)
    developed["quality_profile"] = "professional"
    developed["document_type"] = "narrative_review"
    developed["presentation"] = "hybrid"
    developed["table_of_contents"] = True
    developed["references"] = [
        {"key": f"test{i}", "authors": f"Investigador{i}, A.", "year": "2023",
         "title": f"Fuente sintética de integración {i}", "journal": "Revista de prueba",
         "url": f"https://example.org/synthetic/{i}"}
        for i in range(8)
    ]
    chapter_kinds = [
        ("abstract", "Resumen"), ("introduction", "Introducción"),
        ("methods", "Método"), ("results", "Resultados"),
        ("discussion", "Discusión"), ("limitations", "Limitaciones"),
        ("conclusions", "Conclusiones"),
    ]
    developed["sections"] = [
        {"kind": kind, "heading": heading,
         "text": " ".join(
             "Fragmento sintético número " + str(j) + " de " + heading +
             " donde se explican criterios de prueba documentados, comparaciones conceptuales, "
             "restricciones metodológicas y requisitos editoriales verificables para la "
             "integración automática del informe [@test" + str((j+index) % 8) + "]."
             for j in range(22)
         )}
        for index, (kind, heading) in enumerate(chapter_kinds)
    ]
    developed["tables"] = [
        {"title": "Matriz de evidencia exclusivamente sintética",
         "columns": ["Fuente", "Categoría", "Observación"],
         "rows": [[f"test{i}", "Sintética", f"Prueba {i}"] for i in range(8)],
         "source_key": "test0"}
    ]
    developed["charts"] = [
        {"title": "Comparación gráfica de valores solo sintéticos",
         "labels": ["A", "B", "C"], "values": [4, 7, 5],
         "source_key": "test0"}
    ]
    output = subprocess.run([sys.executable, str(script)],
                            input=json.dumps(developed, ensure_ascii=False).encode("utf-8"),
                            capture_output=True, timeout=120)
    assert output.returncode == 0, output.stderr.decode("utf-8", errors="replace")
    complete = json.loads(output.stdout.decode("utf-8"))
    assert complete["state"] == "REPORT_READY", complete
    assert complete["quality_profile"] == "professional"
    assert complete["source_count"] == 8
    with fitz.open(complete["files"]["pdf"]) as publication:
        assert publication.page_count >= 6
        fulltext = "\\n".join(page.get_text() for page in publication)
        assert "Índice general" in fulltext
        assert "Matriz de evidencia" in fulltext
        assert "Comparación gráfica" in fulltext
        assert any(page.get_images() for page in publication), "El PDF debe incluir una figura real."
        assert "[@" not in fulltext
    print("RESEARCH_PUBLISHER_SMOKE_PASS")
finally:
    logo.unlink(missing_ok=True)
